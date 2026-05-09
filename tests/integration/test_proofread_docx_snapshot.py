"""端到端：cassette回放 + docx结构化摘要 + syrupy快照。

这一文件示范了"完整管线"测试的目标形态：

- HTTP层 由 cassette 替换 → 离线、确定性、零 token 消耗
- 输出层 用 ``summarize_docx`` 提取 paragraphs/comments/ins/del → 不受zip ID扰动影响
- 结果用 syrupy 快照 → 算法迭代时 diff 一目了然

录制工作流
==========

1. 设 ``LICLOUD_TOKEN`` （或 ``OPENAI_API_KEY``、``ANTHROPIC_API_KEY``）
2. ``rm tests/cassettes/Test*proofread*``  + ``rm tests/integration/__snapshots__/test_proofread_docx_snapshot.ambr``
3. ``uv run pytest tests/integration/test_proofread_docx_snapshot.py --record-mode=once --snapshot-update``
4. 人工 review cassette 与 snapshot 的 diff，提交
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest
from docx import Document

# 让 tests/utils 可被 import
TESTS_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(TESTS_ROOT))

from utils.docx_diff import normalize_for_snapshot, summarize_docx
from utils.docx_lint import lint_docx

pytestmark = [pytest.mark.integration, pytest.mark.vcr]


def _make_input_docx(path: Path) -> Path:
    """生成可重复的输入文档（避免依赖 sample_input.docx 的内容变化）。"""
    doc = Document()
    doc.add_paragraph("第1章 计算机基础")
    doc.add_paragraph("操做系统是管理计算机硬件和软件资源的程式。")
    doc.add_paragraph("常见的操作系统包括Windows、Linux和macOS。")
    doc.save(path)
    return path


@pytest.fixture
def licloud_env(monkeypatch: pytest.MonkeyPatch):
    """统一设置 licloud 后端 + 占位 token（cassette 回放下不会真出网）。"""
    monkeypatch.setenv("AI_PROVIDER", "licloud")
    monkeypatch.setenv("LICLOUD_TOKEN", os.getenv("LICLOUD_TOKEN", "test-licloud-token"))
    monkeypatch.setenv("LICLOUD_MODEL", os.getenv("LICLOUD_MODEL", "aws-claude-sonnet-4-6"))


class TestProofreadDocxSnapshot:
    """端到端：input docx → AI → output docx → 摘要快照。"""

    def test_comments_mode_pipeline(self, licloud_env, tmp_path, snapshot):
        from proofreader_factory import create_proofreader

        in_path = _make_input_docx(tmp_path / "input.docx")
        out_path = tmp_path / "output.docx"

        pr = create_proofreader(os.getenv("LICLOUD_TOKEN", "test"), "comments")
        ok = pr.proofread_document(str(in_path), str(out_path))
        assert ok, "校对管线失败"

        # OOXML lint —— 钉死 Word "无法读取的内容/恢复" 的几类已知坑
        lint_issues = lint_docx(out_path)
        errors = [i for i in lint_issues if i.severity == "error"]
        assert not errors, "docx OOXML 不合规（Word 会拒收）:\n" + "\n".join(
            f"  - {e}" for e in errors
        )

        summary = summarize_docx(out_path)
        # 段落原文应该被保留
        assert any("操作系统" in p or "操做系统" in p for p in summary.paragraphs)
        # 输出做规范化后快照
        assert normalize_for_snapshot(summary) == snapshot
