"""End-to-end integration tests for the proofreading pipeline."""

import os
import sys
import tempfile
import pytest
from unittest.mock import patch, MagicMock
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'src'))


@pytest.mark.integration
class TestEndToEnd:
    """Integration tests that exercise the full pipeline with mocked AI."""

    def _create_sample_docx(self, path: str) -> str:
        """Helper: create a .docx file with Chinese text and return its path."""
        doc = Document()
        doc.add_paragraph("第1章 计算机基础知识")
        doc.add_paragraph("计算机是一种能够按照程序运行的电子设备。")
        doc.add_paragraph("操做系统是管理计算机硬件和软件资源的程序。")
        doc.add_paragraph("常见的操作系统包括Windows、Linux和macOS。")
        doc.save(path)
        return path

    # ------------------------------------------------------------------
    # comments mode
    # ------------------------------------------------------------------

    @patch("ai_checker.openai.OpenAI")
    def test_full_pipeline_comments_mode(self, mock_openai_cls):
        """Run the full pipeline in comments mode with mocked AI responses."""
        from proofreader_factory import create_proofreader

        mock_ai_result = {
            "issues": [
                {
                    "type": "错别字",
                    "text": "操做",
                    "suggestion": "建议改为：操作",
                    "severity": "high",
                },
            ],
            "suggestions": [],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "sample.docx")
            output_path = os.path.join(tmpdir, "sample_comments.docx")
            self._create_sample_docx(input_path)

            pr = create_proofreader("fake-key", "comments")

            with patch.object(pr.ai_checker, "_ai_proofread", return_value=mock_ai_result), \
                 patch.object(pr.ai_checker, "_check_terminology", return_value=[]):
                result = pr.proofread_document(input_path, output_path)

            assert result is True
            assert os.path.exists(output_path)
            # Verify the output is a valid docx
            out_doc = Document(output_path)
            assert len(out_doc.paragraphs) > 0

    # ------------------------------------------------------------------
    # revisions mode
    # ------------------------------------------------------------------

    @patch("ai_checker.openai.OpenAI")
    def test_full_pipeline_revisions_mode(self, mock_openai_cls):
        """Run the full pipeline in revisions mode with mocked AI responses."""
        from proofreader_factory import create_proofreader

        mock_ai_result = {
            "issues": [
                {
                    "type": "错别字",
                    "text": "操做",
                    "suggestion": "建议改为：操作",
                    "severity": "high",
                },
                {
                    "type": "用词不当",
                    "text": "电子设备",
                    "suggestion": "建议改为：电子装置",
                    "severity": "low",
                },
            ],
            "suggestions": [],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "sample.docx")
            output_path = os.path.join(tmpdir, "sample_revisions.docx")
            self._create_sample_docx(input_path)

            pr = create_proofreader("fake-key", "revisions")

            with patch.object(pr.ai_checker, "_ai_proofread", return_value=mock_ai_result), \
                 patch.object(pr.ai_checker, "_check_terminology", return_value=[]):
                result = pr.proofread_document(input_path, output_path)

            assert result is True
            assert os.path.exists(output_path)
            out_doc = Document(output_path)
            assert len(out_doc.paragraphs) > 0

    # ------------------------------------------------------------------
    # no issues – output should still be created (copy of original)
    # ------------------------------------------------------------------

    @patch("ai_checker.openai.OpenAI")
    def test_full_pipeline_no_issues(self, mock_openai_cls):
        """When AI returns no issues, the output file is a copy of the original."""
        from proofreader_factory import create_proofreader

        mock_ai_result = {
            "issues": [],
            "suggestions": [],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "sample.docx")
            output_path = os.path.join(tmpdir, "sample_out.docx")
            self._create_sample_docx(input_path)

            pr = create_proofreader("fake-key", "comments")

            with patch.object(pr.ai_checker, "_ai_proofread", return_value=mock_ai_result), \
                 patch.object(pr.ai_checker, "_check_terminology", return_value=[]):
                result = pr.proofread_document(input_path, output_path)

            assert result is True
            assert os.path.exists(output_path)
            # The output should be a valid docx with the same content
            out_doc = Document(output_path)
            texts = [p.text for p in out_doc.paragraphs]
            assert "计算机是一种能够按照程序运行的电子设备。" in texts
