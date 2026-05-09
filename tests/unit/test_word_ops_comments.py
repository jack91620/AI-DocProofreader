"""word_ops.comments模块单元测试"""

import os
import sys
import tempfile
import xml.etree.ElementTree as ET

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))

from docx import Document

from word_ops.comments import (
    WordCommentsManager,
    create_comments_xml,
    create_document_rels,
    update_content_types,
)

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"


class TestWordCommentsManager:
    """WordCommentsManager批注管理器测试"""

    def _make_doc_with_paragraph(self, text="这是一段测试文本"):
        doc = Document()
        para = doc.add_paragraph(text)
        return doc, para

    def test_add_comment_success(self):
        """添加批注成功时返回True且计数器递增"""
        doc, para = self._make_doc_with_paragraph("这是一段测试文本")
        mgr = WordCommentsManager(doc)

        result = mgr.add_comment(para, "测试", "这里有错误")

        assert result is True
        assert mgr.comment_counter == 1
        assert len(mgr.comments) == 1
        assert mgr.comments[0]["id"] == 1
        assert mgr.comments[0]["text"] == "这里有错误"
        assert mgr.comments[0]["author"] == "AI校对助手"

    def test_add_comment_increments_counter(self):
        """多次添加批注时计数器正确递增"""
        doc, para = self._make_doc_with_paragraph("第一处错误和第二处错误")
        mgr = WordCommentsManager(doc)

        mgr.add_comment(para, "第一处", "批注1")
        mgr.add_comment(para, "第二处", "批注2")

        assert mgr.comment_counter == 2
        assert len(mgr.comments) == 2
        assert mgr.comments[0]["id"] == 1
        assert mgr.comments[1]["id"] == 2

    def test_add_comment_text_not_found(self):
        """目标文本不在段落中时返回False"""
        doc, para = self._make_doc_with_paragraph("这是一段测试文本")
        mgr = WordCommentsManager(doc)

        result = mgr.add_comment(para, "不存在的文本", "批注")

        assert result is False
        assert mgr.comment_counter == 0
        assert len(mgr.comments) == 0

    def test_finalize_no_comments(self):
        """没有批注时finalize返回True"""
        doc = Document()
        mgr = WordCommentsManager(doc)

        result = mgr.finalize_document()

        assert result is True

    def test_finalize_with_comments(self):
        """有批注时finalize也返回True"""
        doc, para = self._make_doc_with_paragraph("测试文本")
        mgr = WordCommentsManager(doc)
        mgr.add_comment(para, "测试", "批注内容")

        result = mgr.finalize_document()

        assert result is True

    def test_get_comments_for_xml(self):
        """验证get_comments_for_xml返回正确结构"""
        doc, para = self._make_doc_with_paragraph("第一处和第二处文本")
        mgr = WordCommentsManager(doc)
        mgr.add_comment(para, "第一处", "批注A", author="TestAuthor")
        mgr.add_comment(para, "第二处", "批注B")

        data = mgr.get_comments_for_xml()

        assert len(data) == 2

        assert data[0]["id"] == 1
        assert data[0]["text"] == "批注A"
        assert data[0]["author"] == "TestAuthor"
        assert "date" in data[0]

        assert data[1]["id"] == 2
        assert data[1]["text"] == "批注B"
        assert data[1]["author"] == "AI校对助手"

    def test_get_comments_for_xml_empty(self):
        """没有批注时返回空列表"""
        doc = Document()
        mgr = WordCommentsManager(doc)

        data = mgr.get_comments_for_xml()

        assert data == []


class TestCreateCommentsXml:
    """create_comments_xml函数测试"""

    def _sample_comments(self, count=1):
        return [
            {
                "id": i,
                "text": f"批注内容{i}",
                "author": "AI校对助手",
                "date": "2026-01-01T00:00:00Z",
            }
            for i in range(1, count + 1)
        ]

    def test_creates_valid_xml(self):
        """生成的XML文件存在且包含正确的元素"""
        with tempfile.TemporaryDirectory() as tmp:
            xml_path = os.path.join(tmp, "word", "comments.xml")
            create_comments_xml(xml_path, self._sample_comments(1))

            assert os.path.exists(xml_path)

            tree = ET.parse(xml_path)
            root = tree.getroot()

            # 根元素应是 w:comments
            assert root.tag == f"{{{NS_W}}}comments"

            # 应有一个 w:comment 子元素
            comments = root.findall(f"{{{NS_W}}}comment")
            assert len(comments) == 1

            comment = comments[0]
            assert comment.get(f"{{{NS_W}}}id") == "1"
            assert comment.get(f"{{{NS_W}}}author") == "AI校对助手"

            # comment下应有 w:p > w:r > w:t 结构
            t_elem = comment.find(f"{{{NS_W}}}p/{{{NS_W}}}r/{{{NS_W}}}t")
            assert t_elem is not None
            assert t_elem.text == "批注内容1"

    def test_multiple_comments(self):
        """多个批注时生成多个comment元素"""
        with tempfile.TemporaryDirectory() as tmp:
            xml_path = os.path.join(tmp, "word", "comments.xml")
            create_comments_xml(xml_path, self._sample_comments(3))

            tree = ET.parse(xml_path)
            root = tree.getroot()

            comments = root.findall(f"{{{NS_W}}}comment")
            assert len(comments) == 3

            ids = [c.get(f"{{{NS_W}}}id") for c in comments]
            assert ids == ["1", "2", "3"]


class TestCreateDocumentRels:
    """create_document_rels函数测试"""

    def test_creates_new_rels(self):
        """目录为空时创建新的rels文件并包含comments关系"""
        with tempfile.TemporaryDirectory() as tmp:
            create_document_rels(tmp)

            rels_path = os.path.join(tmp, "word", "_rels", "document.xml.rels")
            assert os.path.exists(rels_path)

            tree = ET.parse(rels_path)
            root = tree.getroot()

            # 查找Relationship元素（不带命名空间前缀，因为注册了默认命名空间）
            all_rels = list(root)
            comment_rels = [
                r
                for r in all_rels
                if r.get("Type")
                == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
            ]
            assert len(comment_rels) == 1
            assert comment_rels[0].get("Target") == "comments.xml"

    def test_updates_existing_rels(self):
        """已存在rels文件时添加关系而不重复"""
        with tempfile.TemporaryDirectory() as tmp:
            rels_dir = os.path.join(tmp, "word", "_rels")
            os.makedirs(rels_dir, exist_ok=True)
            rels_path = os.path.join(rels_dir, "document.xml.rels")

            # 创建一个已有的rels文件（包含一个样式关系）
            existing_xml = (
                '<?xml version="1.0" encoding="utf-8"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                "</Relationships>"
            )
            with open(rels_path, "w", encoding="utf-8") as f:
                f.write(existing_xml)

            create_document_rels(tmp)

            tree = ET.parse(rels_path)
            root = tree.getroot()

            all_rels = list(root)
            comment_rels = [
                r
                for r in all_rels
                if r.get("Type")
                == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
            ]
            assert len(comment_rels) == 1

            # 原有的styles关系应保留
            style_rels = [
                r
                for r in all_rels
                if r.get("Type")
                == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"
            ]
            assert len(style_rels) == 1

    def test_no_duplicate_rels(self):
        """调用两次不产生重复关系"""
        with tempfile.TemporaryDirectory() as tmp:
            create_document_rels(tmp)
            create_document_rels(tmp)

            rels_path = os.path.join(tmp, "word", "_rels", "document.xml.rels")
            tree = ET.parse(rels_path)
            root = tree.getroot()

            all_rels = list(root)
            comment_rels = [
                r
                for r in all_rels
                if r.get("Type")
                == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
            ]
            assert len(comment_rels) == 1


class TestUpdateContentTypes:
    """update_content_types函数测试"""

    def test_creates_content_type(self):
        """目录为空时创建Content_Types文件并添加Override"""
        with tempfile.TemporaryDirectory() as tmp:
            update_content_types(tmp)

            ct_path = os.path.join(tmp, "[Content_Types].xml")
            assert os.path.exists(ct_path)

            tree = ET.parse(ct_path)
            root = tree.getroot()

            all_children = list(root)
            overrides = [e for e in all_children if e.get("PartName") == "/word/comments.xml"]
            assert len(overrides) == 1
            assert (
                overrides[0].get("ContentType")
                == "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
            )

    def test_no_duplicate(self):
        """调用两次只有一个Override"""
        with tempfile.TemporaryDirectory() as tmp:
            update_content_types(tmp)
            update_content_types(tmp)

            ct_path = os.path.join(tmp, "[Content_Types].xml")
            tree = ET.parse(ct_path)
            root = tree.getroot()

            all_children = list(root)
            overrides = [e for e in all_children if e.get("PartName") == "/word/comments.xml"]
            assert len(overrides) == 1


class TestNamespaceCorrectness:
    """钉死回归: Override / Relationship 必须与父节点同命名空间，
    否则 Word 报"无法读取的内容/恢复此文档"。

    历史 bug: ``ET.SubElement(root, "Override")`` 创建无命名空间元素。
    当父节点是 ``<ns0:Types xmlns:ns0="...">`` 这种**前缀绑定**形式时，
    序列化出的 ``<Override>`` 子元素会被解析为"空命名空间"，与父节点
    的 ns0 不同 — Word 拒收。
    """

    NS_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
    NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"

    @staticmethod
    def _ns_of(elem) -> str:
        """从 Clark 表示法 ``{ns}tag`` 提取命名空间，无则返回空串。"""
        tag = elem.tag
        if isinstance(tag, str) and tag.startswith("{"):
            return tag[1 : tag.index("}")]
        return ""

    def test_relationship_has_correct_namespace_in_prefixed_root(self):
        """复现历史 bug 现场: 父节点用 ``xmlns:ns0=`` 前缀绑定。"""
        with tempfile.TemporaryDirectory() as tmp:
            rels_dir = os.path.join(tmp, "word", "_rels")
            os.makedirs(rels_dir, exist_ok=True)
            rels_path = os.path.join(rels_dir, "document.xml.rels")

            # 关键: 用前缀绑定（不是默认 xmlns）—— 还原真实生产环境
            existing = (
                '<?xml version="1.0" encoding="utf-8"?>'
                f'<ns0:Relationships xmlns:ns0="{self.NS_RELS}">'
                '<ns0:Relationship Id="rId1" Type="x" Target="styles.xml"/>'
                "</ns0:Relationships>"
            )
            with open(rels_path, "w", encoding="utf-8") as f:
                f.write(existing)

            create_document_rels(tmp)

            root = ET.parse(rels_path).getroot()
            assert root.tag == f"{{{self.NS_RELS}}}Relationships"
            for child in root:
                ns = self._ns_of(child)
                assert ns == self.NS_RELS, (
                    f"元素 {child.tag} 命名空间为 {ns!r}，应为 {self.NS_RELS}；"
                    "Word 会拒收"
                )

    def test_override_has_correct_namespace_in_prefixed_root(self):
        """复现历史 bug 现场: Content_Types 父节点前缀绑定。"""
        with tempfile.TemporaryDirectory() as tmp:
            ct_path = os.path.join(tmp, "[Content_Types].xml")
            existing = (
                '<?xml version="1.0" encoding="utf-8"?>'
                f'<ns0:Types xmlns:ns0="{self.NS_CT}">'
                '<ns0:Default Extension="xml" ContentType="application/xml"/>'
                "</ns0:Types>"
            )
            with open(ct_path, "w", encoding="utf-8") as f:
                f.write(existing)

            update_content_types(tmp)

            root = ET.parse(ct_path).getroot()
            assert root.tag == f"{{{self.NS_CT}}}Types"
            for child in root:
                ns = self._ns_of(child)
                assert ns == self.NS_CT, (
                    f"元素 {child.tag} 命名空间为 {ns!r}，应为 {self.NS_CT}；"
                    "Word 会拒收"
                )
