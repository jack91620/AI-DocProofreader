"""Word跟踪更改模块单元测试"""

import os
import sys
import tempfile
import xml.etree.ElementTree as ET

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))

from docx import Document
from docx.oxml.ns import qn

from word_ops.track_changes import (
    WordTrackChangesManager,
    enable_track_changes_setting,
)


class TestWordTrackChangesManager:
    """WordTrackChangesManager核心功能测试"""

    def test_add_tracked_change(self):
        """添加跟踪更改后，返回True，revision_counter递增，revisions_data填充"""
        doc = Document()
        para = doc.add_paragraph("这是一个测试句子")
        mgr = WordTrackChangesManager(doc)

        result = mgr.add_tracked_change(para, "测试", "测验", reason="用词修正")

        assert result is True
        assert mgr.revision_counter == 1
        assert len(mgr.revisions_data) == 1

        rev = mgr.revisions_data[0]
        assert rev["original_text"] == "测试"
        assert rev["corrected_text"] == "测验"
        assert rev["reason"] == "用词修正"
        assert rev["del_id"] == 1
        assert rev["ins_id"] == 1001
        assert rev["author"] == "AI校对助手"

    def test_add_multiple_changes_same_paragraph(self):
        """同一段落添加两个更改，paragraph_changes正确跟踪"""
        doc = Document()
        para = doc.add_paragraph("第一个错误和第二个错误在这里")
        mgr = WordTrackChangesManager(doc)

        mgr.add_tracked_change(para, "第一个错误", "第一个正确")
        mgr.add_tracked_change(para, "第二个错误", "第二个正确")

        assert mgr.revision_counter == 2
        assert len(mgr.revisions_data) == 2

        para_id = id(para)
        assert para_id in mgr.paragraph_changes
        assert len(mgr.paragraph_changes[para_id]["changes"]) == 2
        assert mgr.paragraph_changes[para_id]["changes"][0]["original_text"] == "第一个错误"
        assert mgr.paragraph_changes[para_id]["changes"][1]["original_text"] == "第二个错误"

    def test_apply_all_changes(self):
        """应用更改后，段落XML包含w:del和w:ins元素"""
        doc = Document()
        para = doc.add_paragraph("原始文本需要修改")
        mgr = WordTrackChangesManager(doc)

        mgr.add_tracked_change(para, "原始", "修正", reason="纠错")
        mgr.apply_all_changes()

        xml_str = para._element.xml
        assert "w:del" in xml_str
        assert "w:ins" in xml_str

    def test_unique_marker_prefix(self):
        """两个管理器实例的_marker_prefix应不同"""
        doc = Document()
        mgr1 = WordTrackChangesManager(doc)
        mgr2 = WordTrackChangesManager(doc)

        assert mgr1._marker_prefix != mgr2._marker_prefix
        assert len(mgr1._marker_prefix) == 8
        assert len(mgr2._marker_prefix) == 8


class TestCreateElements:
    """删除/插入元素创建测试"""

    def _make_manager(self):
        doc = Document()
        return WordTrackChangesManager(doc)

    def test_create_deletion_element(self):
        """w:del元素包含正确的属性和w:delText子元素"""
        mgr = self._make_manager()
        el = mgr._create_deletion_element("被删除文字", 42)

        assert el.tag == qn("w:del")
        assert el.get(qn("w:id")) == "42"
        assert el.get(qn("w:author")) == "AI校对助手"
        assert el.get(qn("w:date")) == mgr.date

        del_texts = el.findall(f".//{qn('w:delText')}")
        assert len(del_texts) == 1
        assert del_texts[0].text == "被删除文字"

    def test_create_insertion_element(self):
        """w:ins元素包含正确的属性和w:t子元素"""
        mgr = self._make_manager()
        el = mgr._create_insertion_element("插入文字", 99)

        assert el.tag == qn("w:ins")
        assert el.get(qn("w:id")) == "99"
        assert el.get(qn("w:author")) == "AI校对助手"
        assert el.get(qn("w:date")) == mgr.date

        t_elements = el.findall(f".//{qn('w:t')}")
        assert len(t_elements) == 1
        assert t_elements[0].text == "插入文字"


class TestEnableTrackChangesSetting:
    """enable_track_changes_setting功能测试"""

    _NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def _make_minimal_settings_xml(self, word_dir):
        """在word_dir下创建最小settings.xml"""
        content = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            f'<w:settings xmlns:w="{self._NS_W}"></w:settings>'
        )
        settings_path = os.path.join(word_dir, "settings.xml")
        with open(settings_path, "w", encoding="utf-8") as f:
            f.write(content)
        return settings_path

    def test_creates_setting(self):
        """已有settings.xml时，添加trackRevisions元素"""
        with tempfile.TemporaryDirectory() as tmp:
            word_dir = os.path.join(tmp, "word")
            os.makedirs(word_dir)
            settings_path = self._make_minimal_settings_xml(word_dir)

            enable_track_changes_setting(tmp)

            tree = ET.parse(settings_path)
            root = tree.getroot()
            ns = {"w": self._NS_W}
            track = root.find(".//w:trackRevisions", ns)

            assert track is not None
            assert track.get(f"{{{self._NS_W}}}val") == "1"

    def test_creates_setting_when_no_file(self):
        """settings.xml不存在时，函数创建新文件并写入trackRevisions"""
        with tempfile.TemporaryDirectory() as tmp:
            word_dir = os.path.join(tmp, "word")
            os.makedirs(word_dir)
            settings_path = os.path.join(word_dir, "settings.xml")

            # 文件不存在
            assert not os.path.exists(settings_path)

            enable_track_changes_setting(tmp)

            # 函数应创建文件
            assert os.path.exists(settings_path)

            tree = ET.parse(settings_path)
            root = tree.getroot()
            ns = {"w": self._NS_W}
            track = root.find(".//w:trackRevisions", ns)
            assert track is not None
            assert track.get(f"{{{self._NS_W}}}val") == "1"
