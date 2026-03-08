"""ProofReader and proofreader_factory unit tests"""

import os
import sys
import tempfile
import pytest
from unittest.mock import patch, MagicMock, PropertyMock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'src'))


# ---------------------------------------------------------------------------
# TestCreateProofreader  (proofreader_factory.create_proofreader)
# ---------------------------------------------------------------------------

class TestCreateProofreader:
    """Test the factory function create_proofreader."""

    @patch("ai_checker.openai.OpenAI")
    def test_creates_with_comments_mode(self, mock_openai):
        from proofreader_factory import create_proofreader
        pr = create_proofreader("fake-key", "comments")
        assert pr is not None
        assert pr.strategy is not None

    @patch("ai_checker.openai.OpenAI")
    def test_creates_with_revisions_mode(self, mock_openai):
        from proofreader_factory import create_proofreader
        pr = create_proofreader("fake-key", "revisions")
        assert pr is not None
        assert pr.strategy is not None

    @patch("ai_checker.openai.OpenAI")
    def test_creates_with_track_changes_mode(self, mock_openai):
        from proofreader_factory import create_proofreader
        pr = create_proofreader("fake-key", "track_changes")
        assert pr is not None
        assert pr.strategy is not None

    @patch("ai_checker.openai.OpenAI")
    def test_creates_with_enhanced_mode(self, mock_openai):
        from proofreader_factory import create_proofreader
        pr = create_proofreader("fake-key", "enhanced")
        assert pr is not None
        assert pr.strategy is not None

    @patch("ai_checker.openai.OpenAI")
    def test_raises_for_unknown_mode(self, mock_openai):
        from proofreader_factory import create_proofreader
        with pytest.raises(ValueError, match="未知的校对模式"):
            create_proofreader("fake-key", "nonexistent_mode")


# ---------------------------------------------------------------------------
# TestProofReader  (proofreader.ProofReader)
# ---------------------------------------------------------------------------

class TestProofReader:
    """Test the ProofReader class."""

    @patch("ai_checker.openai.OpenAI")
    def test_init_with_strategy(self, mock_openai):
        from proofreader import ProofReader
        mock_strategy = MagicMock()
        pr = ProofReader(api_key="fake-key", strategy=mock_strategy)
        assert pr.strategy is mock_strategy

    @patch("ai_checker.openai.OpenAI")
    def test_init_without_strategy(self, mock_openai):
        from proofreader import ProofReader
        pr = ProofReader(api_key="fake-key")
        assert pr.strategy is None

    @patch("ai_checker.openai.OpenAI")
    def test_proofread_document_no_corrections(self, mock_openai):
        """When AI finds no issues the file is copied and True is returned.
        The strategy should NOT be called because there are no corrections."""
        from proofreader import ProofReader
        from ai_checker import ProofreadingResult
        from docx import Document

        # Create a real .docx input file
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "input.docx")
            output_path = os.path.join(tmpdir, "output.docx")

            doc = Document()
            doc.add_paragraph("这是一段测试文本。")
            doc.save(input_path)

            # Build an empty ProofreadingResult (no issues)
            empty_result = ProofreadingResult()

            mock_strategy = MagicMock()
            pr = ProofReader(api_key="fake-key", strategy=mock_strategy)

            with patch.object(pr.ai_checker, "check_text_by_paragraphs", return_value=empty_result):
                result = pr.proofread_document(input_path, output_path)

            assert result is True
            assert os.path.exists(output_path)
            # Strategy should not have been called — no corrections
            mock_strategy.apply_corrections.assert_not_called()

    @patch("ai_checker.openai.OpenAI")
    def test_proofread_document_with_corrections(self, mock_openai):
        """When AI finds issues, strategy.apply_corrections should be called."""
        from proofreader import ProofReader
        from ai_checker import ProofreadingResult
        from docx import Document

        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "input.docx")
            output_path = os.path.join(tmpdir, "output.docx")

            doc = Document()
            doc.add_paragraph("这是一段包含操做错误的测试文本。")
            doc.save(input_path)

            # Build a ProofreadingResult with one issue
            ai_result = ProofreadingResult()
            ai_result.add_issue(
                issue_type="错别字",
                text="操做",
                suggestion="建议改为：操作",
                severity="high",
            )

            mock_strategy = MagicMock()
            mock_strategy.apply_corrections.return_value = True
            mock_strategy.get_output_suffix.return_value = "proofread"

            pr = ProofReader(api_key="fake-key", strategy=mock_strategy)

            with patch.object(pr.ai_checker, "check_text_by_paragraphs", return_value=ai_result):
                result = pr.proofread_document(input_path, output_path)

            assert result is True
            mock_strategy.apply_corrections.assert_called_once()
