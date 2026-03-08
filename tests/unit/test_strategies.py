"""输出策略模块单元测试"""

import os
import sys
import tempfile
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'src'))

from docx import Document
from strategies.base import Correction, OutputStrategy
from strategies.comments import CommentsStrategy
from strategies.track_changes import TrackChangesStrategy
from strategies.enhanced import EnhancedStrategy
from strategies.visual_revisions import VisualRevisionsStrategy


class TestCorrection:
    """Correction数据类测试"""

    def test_creation(self):
        c = Correction(
            paragraph_index=2,
            original_text="原始文本",
            corrected_text="修正文本",
            reason="语法错误",
            correction_type="issue",
            severity="high",
            chunk_info={"start": 0, "end": 100},
        )
        assert c.paragraph_index == 2
        assert c.original_text == "原始文本"
        assert c.corrected_text == "修正文本"
        assert c.reason == "语法错误"
        assert c.correction_type == "issue"
        assert c.severity == "high"
        assert c.chunk_info == {"start": 0, "end": 100}

    def test_default_severity(self):
        c = Correction(
            paragraph_index=0,
            original_text="a",
            corrected_text="b",
            reason="r",
            correction_type="issue",
        )
        assert c.severity == "medium"

    def test_optional_chunk_info(self):
        c = Correction(
            paragraph_index=0,
            original_text="a",
            corrected_text="b",
            reason="r",
            correction_type="issue",
        )
        assert c.chunk_info is None


class TestOutputStrategySuffixes:
    """各策略输出后缀测试"""

    def test_comments_suffix(self):
        assert CommentsStrategy().get_output_suffix() == "proofread"

    def test_track_changes_suffix(self):
        assert TrackChangesStrategy().get_output_suffix() == "tracked"

    def test_enhanced_suffix(self):
        assert EnhancedStrategy().get_output_suffix() == "tracked_with_comments_fixed"

    def test_visual_revisions_suffix(self):
        assert VisualRevisionsStrategy().get_output_suffix() == "revised"


def _make_test_doc(paragraphs_text):
    """创建包含指定段落的测试文档，保存到临时文件并返回 (doc, path)"""
    doc = Document()
    for text in paragraphs_text:
        doc.add_paragraph(text)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)
    return doc, tmp.name


class TestVisualRevisionsStrategy:
    """视觉修订策略测试"""

    def test_apply_corrections_success(self):
        doc, input_path = _make_test_doc(["这是一个测试段落", "第二个段落内容"])
        output_path = input_path.replace(".docx", "_output.docx")
        try:
            corrections = [
                Correction(
                    paragraph_index=0,
                    original_text="测试",
                    corrected_text="测验",
                    reason="用词修正",
                    correction_type="issue",
                ),
            ]
            strategy = VisualRevisionsStrategy()
            result = strategy.apply_corrections(doc, corrections, ["这是一个测试段落", "第二个段落内容"], input_path, output_path)
            assert result is True
            assert os.path.exists(output_path)
        finally:
            for p in [input_path, output_path]:
                if os.path.exists(p):
                    os.unlink(p)

    def test_apply_corrections_empty(self):
        doc, input_path = _make_test_doc(["一段文字"])
        output_path = input_path.replace(".docx", "_output.docx")
        try:
            strategy = VisualRevisionsStrategy()
            result = strategy.apply_corrections(doc, [], ["一段文字"], input_path, output_path)
            assert result is True
            assert os.path.exists(output_path)
        finally:
            for p in [input_path, output_path]:
                if os.path.exists(p):
                    os.unlink(p)

    def test_skip_identical_corrections(self):
        doc, input_path = _make_test_doc(["不需要修改的段落"])
        output_path = input_path.replace(".docx", "_output.docx")
        try:
            corrections = [
                Correction(
                    paragraph_index=0,
                    original_text="不需要修改的段落",
                    corrected_text="不需要修改的段落",  # same as original
                    reason="无需修改",
                    correction_type="issue",
                ),
            ]
            strategy = VisualRevisionsStrategy()
            result = strategy.apply_corrections(doc, corrections, ["不需要修改的段落"], input_path, output_path)
            assert result is True
            # Verify the paragraph text is unchanged (no strikethrough/underline applied)
            saved_doc = Document(output_path)
            assert saved_doc.paragraphs[0].text == "不需要修改的段落"
        finally:
            for p in [input_path, output_path]:
                if os.path.exists(p):
                    os.unlink(p)


class TestCommentsStrategy:
    """批注模式策略测试"""

    def test_apply_corrections_creates_output(self):
        doc, input_path = _make_test_doc(["计算机科学导论", "本章介绍基本概念"])
        # Use a temp directory so _temp.docx intermediate files are writable
        tmpdir = os.path.dirname(input_path)
        output_path = os.path.join(tmpdir, "comments_output.docx")
        try:
            corrections = [
                Correction(
                    paragraph_index=0,
                    original_text="导论",
                    corrected_text="概论",
                    reason="术语统一",
                    correction_type="terminology",
                ),
            ]
            strategy = CommentsStrategy()
            result = strategy.apply_corrections(doc, corrections, ["计算机科学导论", "本章介绍基本概念"], input_path, output_path)
            assert result is True
            assert os.path.exists(output_path)
        finally:
            for p in [input_path, output_path]:
                if os.path.exists(p):
                    os.unlink(p)
            # Clean up possible temp file
            temp_file = output_path.replace(".docx", "_temp.docx")
            if os.path.exists(temp_file):
                os.unlink(temp_file)


class TestTrackChangesStrategy:
    """跟踪更改策略测试"""

    def test_apply_corrections_creates_output(self):
        doc, input_path = _make_test_doc(["这是原始文本内容", "第二段内容"])
        tmpdir = os.path.dirname(input_path)
        output_path = os.path.join(tmpdir, "tracked_output.docx")
        try:
            corrections = [
                Correction(
                    paragraph_index=0,
                    original_text="原始",
                    corrected_text="修改后",
                    reason="内容修正",
                    correction_type="issue",
                ),
            ]
            strategy = TrackChangesStrategy()
            result = strategy.apply_corrections(doc, corrections, ["这是原始文本内容", "第二段内容"], input_path, output_path)
            assert result is True
            assert os.path.exists(output_path)
        finally:
            for p in [input_path, output_path]:
                if os.path.exists(p):
                    os.unlink(p)
            temp_file = output_path.replace(".docx", "_temp.docx")
            if os.path.exists(temp_file):
                os.unlink(temp_file)
