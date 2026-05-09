"""
增强模式策略 - 跟踪更改+批注同时显示
"""

import os
import re
from datetime import datetime

from docx import Document
from rich.console import Console

from word_ops.comments import (
    create_comments_xml,
    create_document_rels,
    update_content_types,
)
from word_ops.docx_packager import modify_docx_package
from word_ops.track_changes import WordTrackChangesManager

from .base import Correction, OutputStrategy


class EnhancedStrategy(OutputStrategy):
    """增强模式输出策略 - 跟踪更改+批注"""

    def __init__(self):
        self.console = Console()

    def get_output_suffix(self) -> str:
        return "tracked_with_comments_fixed"

    def apply_corrections(
        self,
        doc: Document,
        corrections: list[Correction],
        text_content: list[str],
        input_file: str,
        output_file: str,
    ) -> bool:
        try:
            if not corrections:
                self.console.print("[yellow]⚠️ 未发现需要修改的问题，直接复制原文档[/yellow]")
                import shutil

                shutil.copy2(input_file, output_file)
                return True

            # 创建同步更改数据
            synchronized_changes = self._build_synchronized_changes(corrections, text_content, doc)
            self.console.print(f"[green]✅ 创建了 {len(synchronized_changes)} 个同步更改[/green]")

            # 应用同步更改
            return self._apply_synchronized_changes(doc, synchronized_changes, output_file)

        except Exception as e:
            self.console.print(f"[red]❌ 增强模式校对失败: {e}[/red]")
            import traceback

            traceback.print_exc()
            return False

    def _build_synchronized_changes(
        self, corrections: list[Correction], text_content: list[str], doc: Document
    ) -> list:
        """从Correction列表构建同步更改数据"""
        synchronized_changes = []

        # 创建段落索引映射 - 从非空段落索引到实际段落索引
        paragraph_mapping = {}
        text_index = 0
        for doc_index, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                paragraph_mapping[text_index] = doc_index
                text_index += 1

        for correction in corrections:
            if (
                not correction.corrected_text
                or correction.corrected_text == correction.original_text
            ):
                continue

            # 在所有段落中查找匹配项
            for i, paragraph_text in enumerate(text_content):
                if correction.original_text in paragraph_text:
                    actual_paragraph_index = paragraph_mapping.get(i, i)
                    occurrences = paragraph_text.count(correction.original_text)

                    for occurrence in range(occurrences):
                        # 创建批注文本
                        if correction.correction_type == "suggestion":
                            comment_text = f"💡 改进建议: {correction.original_text} → {correction.corrected_text}\n"
                            comment_text += f"📋 原因: {correction.reason}\n"
                            comment_text += "🎯 类型: 改进建议\n"
                        elif correction.correction_type == "terminology":
                            comment_text = f"🔍 术语不一致修正: {correction.original_text} → {correction.corrected_text}\n"
                            comment_text += f"📝 理由: {correction.reason}\n"
                        else:
                            comment_text = f"🔧 错误修正: {correction.original_text} → {correction.corrected_text}\n"
                            comment_text += f"📝 理由: {correction.reason}\n"

                        comment_text += f"⏰ 处理时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

                        synchronized_changes.append(
                            {
                                "paragraph_index": actual_paragraph_index,
                                "original_text": correction.original_text,
                                "corrected_text": correction.corrected_text,
                                "comment_text": comment_text,
                                "reason": correction.reason,
                                "type": correction.correction_type,
                                "occurrence_index": occurrence,
                            }
                        )

        return synchronized_changes

    def _apply_synchronized_changes(
        self, doc: Document, synchronized_changes: list, output_file: str
    ) -> bool:
        """应用同步更改"""
        try:
            track_changes_manager = WordTrackChangesManager(doc)

            applied_count = 0
            comments_data = []

            for change in synchronized_changes:
                paragraph_index = change.get("paragraph_index", 0)
                original_text = change.get("original_text", "")
                corrected_text = change.get("corrected_text", "")
                comment_text = change.get("comment_text", "")
                reason = change.get("reason", "")

                if paragraph_index < len(doc.paragraphs):
                    paragraph = doc.paragraphs[paragraph_index]

                    track_success = track_changes_manager.add_tracked_change(
                        paragraph, original_text, corrected_text, reason
                    )

                    if track_success:
                        comments_data.append(
                            {
                                "id": applied_count + 1,
                                "text": comment_text,
                                "author": "AI校对助手",
                                "date": datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ"),
                                "paragraph_index": paragraph_index,
                                "target_text": original_text,
                            }
                        )
                        applied_count += 1

            # 应用所有跟踪更改
            track_changes_manager.apply_all_changes()

            # 保存临时文档（只包含跟踪更改）
            temp_file = output_file.replace(".docx", "_temp_with_track_changes.docx")
            doc.save(temp_file)

            # 创建最终文档（带跟踪更改+批注）
            success = self._create_final_document(
                temp_file, output_file, track_changes_manager.revisions_data, comments_data
            )

            # 清理临时文件
            if os.path.exists(temp_file):
                os.remove(temp_file)

            if success:
                self.console.print(f"[green]✅ 成功应用 {applied_count} 个更改和批注[/green]")
            return success

        except Exception as e:
            self.console.print(f"[red]❌ 应用更改和批注失败: {e}[/red]")
            return False

    def _create_final_document(
        self, temp_file: str, output_file: str, revisions_data, comments_data: list
    ) -> bool:
        """创建最终文档（跟踪更改+批注）"""
        try:
            import shutil

            if not comments_data:
                shutil.copy2(temp_file, output_file)
                return True

            def modifier(temp_dir):
                # 1. 启用跟踪更改
                from word_ops.track_changes import (
                    enable_track_changes_setting,
                    verify_document_revisions,
                )

                enable_track_changes_setting(temp_dir)
                verify_document_revisions(temp_dir)

                # 2. 更新文档关系 + 内容类型（先于 comments.xml 写入，保持顺序）
                create_document_rels(temp_dir)
                update_content_types(temp_dir)

                # 3. 添加批注引用到 document.xml（regex 匹配 <w:del w:id=X>）
                document_xml_path = os.path.join(temp_dir, "word", "document.xml")
                with open(document_xml_path, encoding="utf-8") as f:
                    document_content = f.read()

                updated_document = self._add_comment_references_to_document(
                    document_content, comments_data
                )
                with open(document_xml_path, "w", encoding="utf-8") as f:
                    f.write(updated_document)

                # 4. 扫描 document.xml 中真正被引用的 comment id，
                #    只把这些 comment 写入 comments.xml；orphan 直接丢弃。
                #    背景: 修 D3 (revision id 全局唯一) 之后，同 original_text
                #    多次匹配会复制 change 并分配新 id，导致 comments_data 里的
                #    id 与 document.xml 里实际 <w:del> 的 id 不再一一对应；
                #    某些 comment_id 在 document.xml 里找不到对应 <w:del>，
                #    就成了 orphan。Word 严格模式下 orphan 会触发"无法读取的内容"。
                referenced_ids = set(
                    re.findall(r'<w:commentReference\s+w:id="(\d+)"', updated_document)
                )
                effective_comments = [
                    c for c in comments_data if str(c.get("id")) in referenced_ids
                ]
                comments_xml_path = os.path.join(temp_dir, "word", "comments.xml")
                create_comments_xml(comments_xml_path, effective_comments)

                return True

            return modify_docx_package(temp_file, output_file, modifier)

        except Exception as e:
            self.console.print(f"[red]❌ 创建最终文档失败: {e}[/red]")
            try:
                import shutil

                shutil.copy2(temp_file, output_file)
                return True
            except Exception:
                return False

    def _add_comment_references_to_document(self, document_xml: str, comments_data: list) -> str:
        """在document.xml中添加批注引用标记"""
        try:
            modified_xml = document_xml

            for comment in comments_data:
                comment_id = comment["id"]

                del_pattern = f'(<w:del[^>]*w:id="{comment_id}"[^>]*>.*?</w:del>)'
                matches = re.finditer(del_pattern, modified_xml, re.DOTALL)

                for match in matches:
                    comment_range_start = f'<w:commentRangeStart w:id="{comment_id}"/>'
                    comment_range_end = f'<w:commentRangeEnd w:id="{comment_id}"/>'
                    comment_reference = f'<w:r><w:commentReference w:id="{comment_id}"/></w:r>'

                    replacement = f"{comment_range_start}{match.group(0)}{comment_range_end}{comment_reference}"
                    modified_xml = modified_xml.replace(match.group(0), replacement, 1)
                    break

            return modified_xml

        except Exception as e:
            self.console.print(f"[red]添加批注引用失败: {e}[/red]")
            return document_xml
