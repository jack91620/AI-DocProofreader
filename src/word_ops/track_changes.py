#!/usr/bin/env python
"""
Word跟踪更改操作模块 - 迁移自word_track_changes.py
实现Microsoft Word审阅中的修订功能
"""

import os
import re
import uuid
import xml.etree.ElementTree as ET
from datetime import datetime

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from logging_config import get_logger

from .docx_packager import modify_docx_package

logger = get_logger("word_ops.track_changes")


class WordTrackChangesManager:
    """Word跟踪更改管理器 - 生成真正的Word修订标记"""

    def __init__(self, document):
        self.document = document
        self.revision_counter = 0
        self.author = "AI校对助手"
        self.date = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")
        self.revisions_data = []
        self.paragraph_changes = {}  # 跟踪每个段落的修改
        # 使用UUID前缀避免标记碰撞
        self._marker_prefix = uuid.uuid4().hex[:8]

    def add_tracked_change(
        self, paragraph, original_text: str, corrected_text: str, reason: str = ""
    ):
        """添加跟踪更改（真正的Word修订）"""
        try:
            # 生成修订ID
            self.revision_counter += 1
            del_revision_id = self.revision_counter
            ins_revision_id = self.revision_counter + 1000

            # 存储修订数据
            self.revisions_data.append(
                {
                    "paragraph": paragraph,
                    "original_text": original_text,
                    "corrected_text": corrected_text,
                    "del_id": del_revision_id,
                    "ins_id": ins_revision_id,
                    "author": self.author,
                    "date": self.date,
                    "reason": reason,
                }
            )

            # 跟踪段落修改，批量处理
            para_id = id(paragraph)
            if para_id not in self.paragraph_changes:
                self.paragraph_changes[para_id] = {
                    "paragraph": paragraph,
                    "original_text": paragraph.text,
                    "changes": [],
                }

            self.paragraph_changes[para_id]["changes"].append(
                {
                    "original_text": original_text,
                    "corrected_text": corrected_text,
                    "del_id": del_revision_id,
                    "ins_id": ins_revision_id,
                    "reason": reason,
                }
            )

            logger.info(f"已添加跟踪更改: {original_text} -> {corrected_text}")
            return True

        except Exception as e:
            logger.error(f"添加跟踪更改失败: {e}")
            return False

    def apply_all_changes(self):
        """应用所有跟踪更改到文档"""
        for _para_id, para_data in self.paragraph_changes.items():
            paragraph = para_data["paragraph"]
            original_full_text = para_data["original_text"]
            changes = para_data["changes"]

            # 应用所有更改到段落
            self._apply_changes_to_paragraph(paragraph, original_full_text, changes)

    def _apply_changes_to_paragraph(self, paragraph, original_text, changes):
        """将多个更改应用到单个段落"""
        try:
            # 收集所有需要更改的位置
            changes_with_pos = []
            processed_positions = set()

            for change in changes:
                original_part = change["original_text"]
                start = 0
                first_match = True
                while True:
                    pos = original_text.find(original_part, start)
                    if pos == -1:
                        break

                    position_key = (pos, original_part)
                    if position_key not in processed_positions:
                        # 关键: 同一个 original_part 在段落出现多次时，
                        # 必须为每次出现分配独立的 revision id，否则
                        # <w:del> / <w:ins> 在文档内会出现重复 w:id，
                        # OOXML 严格要求 revision id 全局唯一，Word 拒收。
                        if first_match:
                            change_to_use = change
                            first_match = False
                        else:
                            self.revision_counter += 1
                            change_to_use = {
                                **change,
                                "del_id": self.revision_counter,
                                "ins_id": self.revision_counter + 1000,
                            }
                        changes_with_pos.append((pos, change_to_use))
                        processed_positions.add(position_key)

                    start = pos + 1

            # 按位置降序排序（从后往前处理）
            changes_with_pos.sort(key=lambda x: x[0], reverse=True)

            # 清空段落
            paragraph.clear()

            if not changes_with_pos:
                paragraph.add_run(original_text)
                return

            # 构建标记文本 - 使用UUID前缀避免碰撞
            current_text = original_text
            pfx = self._marker_prefix

            for pos, change in changes_with_pos:
                original_part = change["original_text"]
                corrected_part = change["corrected_text"]
                del_id = change["del_id"]
                ins_id = change["ins_id"]

                if (
                    pos + len(original_part) <= len(current_text)
                    and current_text[pos : pos + len(original_part)] == original_part
                ):
                    before_part = current_text[:pos]
                    after_part = current_text[pos + len(original_part) :]

                    marker = f"\x00{pfx}DEL:{del_id}:{original_part}\x00\x00{pfx}INS:{ins_id}:{corrected_part}\x00"
                    current_text = before_part + marker + after_part

            # 解析并构建段落
            self._build_paragraph_from_marked_text(paragraph, current_text)

        except Exception as e:
            logger.error(f"应用段落更改失败: {e}")
            paragraph.clear()
            paragraph.add_run(original_text)

    def _build_paragraph_from_marked_text(self, paragraph, marked_text):
        """从标记文本构建段落"""
        try:
            pfx = self._marker_prefix
            del_marker = f"\x00{pfx}DEL:"
            ins_marker = f"\x00{pfx}INS:"

            pos = 0
            while pos < len(marked_text):
                del_start = marked_text.find(del_marker, pos)
                if del_start == -1:
                    remaining = marked_text[pos:]
                    if remaining:
                        paragraph.add_run(remaining)
                    break

                # 添加删除标记之前的文本
                if del_start > pos:
                    before_text = marked_text[pos:del_start]
                    paragraph.add_run(before_text)

                # 解析删除标记
                del_content_start = del_start + len(del_marker)
                del_end = marked_text.find("\x00", del_content_start)
                if del_end == -1:
                    break

                del_content = marked_text[del_content_start:del_end]
                parts = del_content.split(":", 1)
                if len(parts) == 2:
                    del_id = parts[0]
                    del_text = parts[1]

                    del_element = self._create_deletion_element(del_text, del_id)
                    paragraph._element.append(del_element)

                # 查找对应的插入标记
                ins_start = marked_text.find(ins_marker, del_end)
                if ins_start == del_end + 1:  # 紧接着的插入标记
                    ins_content_start = ins_start + len(ins_marker)
                    ins_end = marked_text.find("\x00", ins_content_start)
                    if ins_end != -1:
                        ins_content = marked_text[ins_content_start:ins_end]
                        parts = ins_content.split(":", 1)
                        if len(parts) == 2:
                            ins_id = parts[0]
                            ins_text = parts[1]

                            ins_element = self._create_insertion_element(ins_text, ins_id)
                            paragraph._element.append(ins_element)

                        pos = ins_end + 1
                    else:
                        pos = del_end + 1
                else:
                    pos = del_end + 1

        except Exception as e:
            logger.error(f"构建段落失败: {e}")
            paragraph.clear()
            # Clean up markers from text
            clean_text = marked_text
            for char in ["\x00"]:
                clean_text = clean_text.replace(char, "")
            paragraph.add_run(clean_text)

    def _create_deletion_element(self, text, revision_id):
        """创建删除元素 (w:del)"""
        del_element = OxmlElement("w:del")
        del_element.set(qn("w:id"), str(revision_id))
        del_element.set(qn("w:author"), self.author)
        del_element.set(qn("w:date"), self.date)

        del_run = OxmlElement("w:r")

        del_text = OxmlElement("w:delText")
        del_text.text = text
        del_run.append(del_text)

        del_element.append(del_run)
        return del_element

    def _create_insertion_element(self, text, revision_id):
        """创建插入元素 (w:ins)"""
        ins_element = OxmlElement("w:ins")
        ins_element.set(qn("w:id"), str(revision_id))
        ins_element.set(qn("w:author"), self.author)
        ins_element.set(qn("w:date"), self.date)

        ins_run = OxmlElement("w:r")

        ins_text = OxmlElement("w:t")
        ins_text.text = text
        ins_run.append(ins_text)

        ins_element.append(ins_run)
        return ins_element


def enable_track_changes_in_docx(docx_path, output_path, revisions_data):
    """在Word文档中启用跟踪更改并添加修订"""

    def modifier(temp_dir):
        enable_track_changes_setting(temp_dir)
        verify_document_revisions(temp_dir)
        return True

    return modify_docx_package(docx_path, output_path, modifier)


def enable_track_changes_setting(temp_dir):
    """在settings.xml中启用跟踪更改。

    实现注意:
    --------
    settings.xml 的根节点通常带 ``mc:Ignorable="w14"`` 这种字符串引用，
    引用的 ``w14`` 等前缀在 root 上以 ``xmlns:w14=...`` 形式声明。
    用 ``ET.parse`` 解析再 ``tree.write`` 重写会把所有前缀重映射为
    ``ns0/ns1/...``，导致 ``Ignorable="w14"`` 引用了**未声明**的前缀，
    OOXML markup-compatibility (Part 3) 严格校验会让 Word 拒收，弹
    "无法读取的内容/恢复此文档"。

    解法: 用字符串级修改插入 ``<w:trackRevisions w:val="1"/>``，
    完全不动原命名空间声明。
    """
    settings_path = os.path.join(temp_dir, "word", "settings.xml")
    try:
        if not os.path.exists(settings_path):
            # 没文件时退回到全新创建（这种场景不会有 Ignorable 问题）
            ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            root = ET.Element(f"{{{ns_w}}}settings")
            ET.SubElement(root, f"{{{ns_w}}}trackRevisions").set(f"{{{ns_w}}}val", "1")
            ET.ElementTree(root).write(settings_path, encoding="utf-8", xml_declaration=True)
            logger.info("已新建 settings.xml 并启用跟踪更改")
            return

        with open(settings_path, encoding="utf-8") as f:
            text = f.read()

        # 已经启用过就跳过
        if re.search(r"<w:trackRevisions[\s/>]", text):
            logger.info("settings.xml 已含 trackRevisions，跳过")
            return

        # 在 <w:settings ...> 起始标签之后插入 <w:trackRevisions w:val="1"/>
        # 兼容自闭合 <w:settings/> 与展开标签 <w:settings>...</w:settings>
        m = re.search(r"<w:settings\b[^>]*?(/?>)", text)
        if not m:
            logger.warning("未识别 <w:settings> 起始标签，settings.xml 未修改")
            return

        if m.group(1) == "/>":
            # 自闭合 → 展开
            new_open = m.group(0)[:-2] + "><w:trackRevisions w:val=\"1\"/></w:settings>"
            text = text[: m.start()] + new_open + text[m.end():]
        else:
            text = text[: m.end()] + "<w:trackRevisions w:val=\"1\"/>" + text[m.end():]

        with open(settings_path, "w", encoding="utf-8") as f:
            f.write(text)
        logger.info("已在 settings.xml 中启用跟踪更改（字符串级插入，保留原命名空间）")

    except Exception as e:
        logger.error(f"启用跟踪更改设置失败: {e}")


def verify_document_revisions(temp_dir):
    """验证document.xml中的修订标记"""
    try:
        document_path = os.path.join(temp_dir, "word", "document.xml")

        if os.path.exists(document_path):
            with open(document_path, encoding="utf-8") as f:
                content = f.read()

            del_count = content.count("<w:del ")
            ins_count = content.count("<w:ins ")
            deltext_count = content.count("<w:delText>")

            logger.info("document.xml修订标记统计:")
            logger.info(f"   - w:del (删除标记): {del_count}")
            logger.info(f"   - w:ins (插入标记): {ins_count}")
            logger.info(f"   - w:delText (删除文本): {deltext_count}")

            if del_count > 0 or ins_count > 0:
                logger.info("发现Word修订标记")
            else:
                logger.warning("未发现Word修订标记")

    except Exception as e:
        logger.error(f"验证修订标记失败: {e}")
